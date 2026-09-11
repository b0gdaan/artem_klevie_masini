"""Build the thesis template and the corrected dispozicija as DOCX files.

The thesis layout follows the diploma theses defended at Višja strokovna šola Academia Maribor in 2026
(see docs/OBLIKOVANJE_ACADEMIA.md): A4, margins 3/2/2.5/2.5 cm, Times New Roman 12 pt, 1.5 line
spacing, justified text, chapter headings 16 pt bold, subheadings 14 pt and 12 pt bold italic, captions
10 pt centred, page numbers centred from the introduction onwards.

Usage: python thesis/tools/build_docs.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

THESIS_DIR = Path(__file__).resolve().parents[1]
BLACK = RGBColor(0, 0, 0)
GUIDE = RGBColor(0x59, 0x59, 0x59)
JUSTIFY, CENTER, LEFT = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT

STUDENT = "Artem Rakhmanov"
TITLE = "SEO-optimizacija spletnih mest: metode in strategije doseganja visokih pozicij v iskalnikih"
TITLE_EN = "Search Engine Optimization of Websites: Methods and Strategies for Achieving High Search Rankings"
MENTOR = "mag. Dušan Brglez"

# --------------------------------------------------------------------------- shared content

RESEARCH_QUESTIONS = [
    "RV1: Katere skupine SEO-ukrepov (tehnični SEO, vsebinska optimizacija, lokalni SEO, dostopnost) najbolj "
    "izboljšajo vidnost strani v Googlu na izbranem spletnem mestu?",
    "RV2: Ali je z brezplačnimi orodji (Google Search Console, Google Analytics, Lighthouse, WAVE) mogoče doseči "
    "in izmeriti rezultate brez najetja specializirane SEO-agencije?",
    "RV3: V kolikšnem času po uvedbi ukrepov se spremembe pokažejo v prikazih, klikih in povprečnem položaju?",
    "RV4: Katere tehnične in vsebinske napake so na izbranem spletnem mestu najpogostejše?",
]
HYPOTHESES = [
    "H1: Tehnična optimizacija (hitrost nalaganja, mobilna prilagodljivost, pravilna struktura URL in kanonične "
    "povezave) izboljša povprečni položaj optimiziranih strani v Googlu (utežen s prikazi) za vsaj 20 % glede na "
    "kontrolne strani, merjeno 6 tednov po dvotedenskem obdobju uvajanja.",
    "H2: Vsebinska optimizacija (ključne besede, meta oznake, strukturirani podatki) poveča organsko kliknost "
    "(CTR), prilagojeno položaju, za vsaj 15 % glede na kontrolne strani v 6 tednih po obdobju uvajanja.",
    "H3: Na izbranem spletnem mestu lokalna optimizacija (Google Business Profile, lokalne ključne besede, "
    "strukturirani podatki LocalBusiness) poveča število klikov iz lokalnih poizvedb bolj kot splošna vsebinska "
    "optimizacija.",
    "H4 – dodatna hipoteza BTEC 2025/2026 »Na človeka usmerjena informatika za dostopne, prilagodljive in povezane "
    "digitalne izkušnje«: Izboljšanje dostopnosti strani po smernicah WCAG 2.1 (alternativna besedila slik, "
    "semantična struktura HTML, oznake obrazcev, kontrast) izboljša povprečni položaj teh strani glede na kontrolne "
    "strani. Google dostopnosti ne navaja kot neposrednega dejavnika uvrščanja, zato je tudi ničelni izid veljaven "
    "rezultat.",
    "H5 (opisna trditev): Gradnja kakovostnih povratnih povezav z relevantnih slovenskih spletnih mest v obdobju "
    "raziskave poveča Domain Authority (Moz) spletnega mesta; cilj iz prvotne dispozicije je vsaj 10 točk. Ker gre "
    "za eno časovno vrsto brez kontrolne skupine in je Domain Authority metrika podjetja Moz, ki je Google ne "
    "uporablja, se H5 obravnava opisno.",
]
PROBLEM = [
    "Diplomsko delo obravnava optimizacijo spletnih mest za iskalnike (SEO – Search Engine Optimization). "
    "Vidnost spletnega mesta v organskih rezultatih iskanja je za mala podjetja pomemben in razmeroma poceni vir "
    "obiskovalcev. Google je v Sloveniji prevladujoči iskalnik: avgusta 2026 je imel približno 93-odstotni tržni "
    "delež (StatCounter, 2026).",
    "Obstoječa literatura SEO-dejavnike večinoma proučuje na velikih vzorcih spletnih mest (npr. Ziakis idr., "
    "2019; Lewandowski idr., 2021) ali jih opisuje v praktičnih priročnikih. Manj je raziskav, ki bi na enem "
    "konkretnem spletnem mestu nadzorovano primerjale učinke posameznih skupin ukrepov – tehničnih, vsebinskih, "
    "lokalnih in ukrepov dostopnosti. Diplomsko delo to vrzel naslavlja s praktičnim SEO-projektom na spletnem "
    "mestu podjetja, pri katerem se učinki merijo glede na kontrolno skupino strani.",
    "Raziskovalni problem: mala podjetja pogosto vlagajo v plačano oglaševanje (PPC), ne da bi sistematično "
    "izkoristila organsko iskanje. Namen naloge je preveriti, ali je z metodičnim pristopom k SEO in z brezplačnimi "
    "orodji mogoče doseči merljivo izboljšanje vidnosti v Googlu in katere skupine ukrepov k temu največ prispevajo.",
]
GOALS = [
    "izvesti SEO-revizijo spletnega mesta in opredeliti najpogostejše tehnične in vsebinske napake,",
    "pripraviti in po skupinah strani uvesti načrt SEO-ukrepov,",
    "izmeriti učinke ukrepov s podatki Google Search Console in jih statistično ovrednotiti,",
    "pripraviti priporočila za podjetje in ponovljivo orodje za merjenje učinkov (izdelek).",
]
METHODS = [
    "V teoretičnem delu bosta uporabljeni metoda deskripcije za opis temeljnih SEO-konceptov in metoda kompilacije "
    "za pregled znanstvene in strokovne literature (znanstveni članki, dokumentacija Google Search Central, HTTP "
    "Archive Web Almanac, smernice WCAG in zakonodaja).",
    "Praktični del je zasnovan kot kvazieksperiment na enem spletnem mestu. Strani se pred uvedbo ukrepov razdelijo "
    "v skupine (tehnična, vsebinska, lokalna optimizacija, dostopnost) in v kontrolno skupino strani, ki se ne "
    "spreminjajo. Učinek se oceni z metodo razlike razlik (difference-in-differences): sprememba na obravnavanih "
    "straneh se primerja s spremembo na kontrolnih straneh v istem obdobju, kar zmanjša vpliv sezonskosti in "
    "posodobitev Googlovega algoritma. Predobdobje traja 6 tednov, sledita 2 tedna uvajanja in 6 tednov poobdobja. "
    "Negotovost se oceni z bootstrapom po straneh; ker se hkrati preverjajo štiri hipoteze, se uporabi Bonferronijev "
    "popravek. Robustnost se preveri s placebo testom (lažni datum uvedbe) in s primerjavo z enostavno primerjavo "
    "pred/po.",
    "Merjenje po hipotezah: H1 – povprečni položaj v Google Search Console in Core Web Vitals (PageSpeed Insights) "
    "pred in po ukrepih; H2 – CTR v Google Search Console, prilagojen položaju; H3 – kliki iz lokalnih poizvedb v "
    "Google Search Console, opisno tudi metrike Google Business Profile; H4 – revizija dostopnosti z orodjema "
    "Lighthouse in WAVE ter povprečni položaj; H5 – mesečno spremljanje Domain Authority (Moz).",
    "Podatki podjetja se obravnavajo zaupno; v nalogi se objavijo le zbirni rezultati.",
]
PRODUCT = (
    "Praktični del bo vseboval: (1) SEO-revizijo spletnega mesta (Screaming Frog SEO Spider, Google Search Console, "
    "lastni pregled HTML), (2) akcijski načrt z dnevnikom uvedenih ukrepov, (3) ponovljivo orodje v programskem "
    "jeziku Python za zbiranje podatkov, revizijo in statistično oceno učinkov, s preverjanjem ponovljivosti in "
    "predstavitveno spletno stranjo (GitHub Pages), ter (4) poročilo s priporočili za podjetje."
)
OUTLINE = [
    ("1 UVOD", ["1.1 Opis področja in opredelitev problema", "1.2 Namen, cilji in osnovne trditve",
                "1.3 Predpostavke in omejitve", "1.4 Uporabljene raziskovalne metode"]),
    ("2 TEORETIČNI DEL – SEO: OSNOVE IN RAZVOJ", [
        "2.1 Delovanje Googlovih algoritmov in dejavniki uvrstitve", "2.2 Tehnični SEO: hitrost, mobilnost in struktura",
        "2.3 Vsebinska optimizacija in ključne besede", "2.4 Gradnja povratnih povezav (off-page SEO)",
        "2.5 Lokalni SEO in Google Business Profile", "2.6 Dostopnost spletnih mest in SEO"]),
    ("3 PRAKTIČNI DEL – SEO-REVIZIJA IN OPTIMIZACIJA SPLETNEGA MESTA", [
        "3.1 Predstavitev izbranega spletnega mesta in začetno stanje", "3.2 Izvedba SEO-revizije",
        "3.3 Implementacija SEO-ukrepov in merjenje rezultatov", "3.4 Dostopnost spletnega mesta (WCAG) in vpliv na SEO",
        "3.5 Analiza rezultatov in robustnost", "3.6 Odgovori na raziskovalna vprašanja in preverjanje hipotez"]),
    ("4 SKLEP", ["4.1 Refleksija raziskovalnega procesa"]),
    ("5 VIRI IN LITERATURA", []),
    ("6 PRILOGE", []),
]
# APA 7 with Slovenian conventions; *text* is set in italics.
REFERENCES = [
    "Bernal, J. L., Cummins, S. in Gasparrini, A. (2017). Interrupted time series regression for the evaluation of "
    "public health interventions: A tutorial. *International Journal of Epidemiology, 46*(1), 348–355. "
    "https://doi.org/10.1093/ije/dyw098",
    "Brin, S. in Page, L. (1998). The anatomy of a large-scale hypertextual Web search engine. *Computer Networks and "
    "ISDN Systems, 30*(1–7), 107–117. https://doi.org/10.1016/S0169-7552(98)00110-X",
    "Cameron, A. C., Gelbach, J. B. in Miller, D. L. (2008). Bootstrap-based improvements for inference with "
    "clustered errors. *The Review of Economics and Statistics, 90*(3), 414–427.",
    "Craswell, N., Zoeter, O., Taylor, M. in Ramsey, B. (2008). An experimental comparison of click position-bias "
    "models. V *Proceedings of the 2008 International Conference on Web Search and Data Mining* (str. 87–94). ACM. "
    "https://doi.org/10.1145/1341531.1341545",
    "Enge, E., Spencer, S. in Stricchiola, J. (2023). *The art of SEO: Mastering search engine optimization* "
    "(4. izd.). O'Reilly Media.",
    "Google. (b. d.-a). *Creating helpful, reliable, people-first content*. Google Search Central. Pridobljeno "
    "[datum] s https://developers.google.com/search/docs/fundamentals/creating-helpful-content",
    "Google. (b. d.-b). *SEO Starter Guide: The basics*. Google Search Central. Pridobljeno [datum] s "
    "https://developers.google.com/search/docs/fundamentals/seo-starter-guide",
    "Google. (b. d.-c). *Understanding Core Web Vitals and Google search results*. Google Search Central. "
    "Pridobljeno [datum] s https://developers.google.com/search/docs/appearance/core-web-vitals",
    "HTTP Archive. (2024). *Web Almanac 2024: SEO*. Pridobljeno [datum] s https://almanac.httparchive.org/en/2024/seo",
    "Lewandowski, D. (2023). *Understanding search engines*. Springer. https://doi.org/10.1007/978-3-031-22789-9",
    "Lewandowski, D., Sünkler, S. in Yagci, N. (2021). The influence of search engine optimization on Google's "
    "results: A multi-dimensional approach for detecting SEO. V *Proceedings of the 13th ACM Web Science Conference "
    "2021*. ACM. https://doi.org/10.1145/3447535.3462479",
    "Matošević, G., Dobša, J. in Mladenić, D. (2021). Using machine learning for web page classification in search "
    "engine optimization. *Future Internet, 13*(1), 9. https://doi.org/10.3390/fi13010009",
    "Moz. (b. d.). *Domain Authority: What is it and how is it calculated*. Pridobljeno [datum] s "
    "https://moz.com/learn/seo/domain-authority",
    "Pearson. (2022). *Pearson BTEC Levels 4 and 5 Higher Nationals in Computing: Specification* (Issue 2). "
    "Pearson Education.",
    "Schultheiß, S. in Lewandowski, D. (2021). \"Outside the industry, nobody knows what we do\": SEO as seen by "
    "search engine optimizers and content providers. *Journal of Documentation, 77*(2), 542–557. "
    "https://doi.org/10.1108/JD-07-2020-0127",
    "StatCounter. (2026). *Search engine market share Slovenia*. Pridobljeno [datum] s "
    "https://gs.statcounter.com/search-engine-market-share/all/slovenia",
    "W3C. (2018). *Web Content Accessibility Guidelines (WCAG) 2.1*. Pridobljeno [datum] s https://www.w3.org/TR/WCAG21/",
    "Whitespark. (2026). *2026 Local search ranking factors*. Pridobljeno [datum] s "
    "https://whitespark.ca/local-search-ranking-factors/",
    "Zakon o dostopnosti do proizvodov in storitev za invalide (ZDPSI). (2023). *Uradni list RS*, št. 14/23.",
    "Ziakis, C., Vlachopoulou, M., Kyrkoudis, T. in Karagkiozidou, M. (2019). Important factors for improving Google "
    "search rank. *Future Internet, 11*(2), 32. https://doi.org/10.3390/fi11020032",
]

# --------------------------------------------------------------------------- low-level helpers


def _clean_theme(rpr):
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is not None:
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)
    color = rpr.find(qn("w:color"))
    if color is not None:
        for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
            color.attrib.pop(qn(attr), None)


def style_font(style, name, size, bold=False, italic=False, caps=False, small_caps=False):
    font = style.font
    font.name, font.size, font.bold, font.italic = name, Pt(size), bold, italic
    font.all_caps, font.small_caps = caps, small_caps
    font.color.rgb = BLACK
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    _clean_theme(rpr)


def paragraph_style(doc, name, base="Normal"):
    try:
        return doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles[base]
        return style


def add_field(paragraph, instruction, placeholder=""):
    def char(kind):
        run = OxmlElement("w:r")
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        run.append(element)
        return run

    p = paragraph._p
    p.append(char("begin"))
    run = OxmlElement("w:r")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = f" {instruction} "
    run.append(text)
    p.append(run)
    p.append(char("separate"))
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.set(qn("xml:space"), "preserve")
    text.text = placeholder
    run.append(text)
    p.append(run)
    p.append(char("end"))


def add_rich(paragraph, text, **run_format):
    for index, part in enumerate(re.split(r"\*", text)):
        if part:
            run = paragraph.add_run(part)
            run.italic = index % 2 == 1 or run_format.get("italic", False)
            if run_format.get("bold"):
                run.bold = True
            if run_format.get("color") is not None:
                run.font.color.rgb = run_format["color"]
            if run_format.get("size"):
                run.font.size = Pt(run_format["size"])
    return paragraph


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:color"), "auto")
    element.set(qn("w:fill"), fill)
    tc_pr.append(element)


def update_fields_on_open(doc):
    element = OxmlElement("w:updateFields")
    element.set(qn("w:val"), "true")
    doc.settings.element.append(element)


def page_setup(section, left, right, top, bottom):
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin, section.right_margin = Cm(left), Cm(right)
    section.top_margin, section.bottom_margin = Cm(top), Cm(bottom)
    section.footer_distance = Cm(1.25)

# --------------------------------------------------------------------------- thesis template


def thesis_styles(doc):
    normal = doc.styles["Normal"]
    style_font(normal, "Times New Roman", 12)
    fmt = normal.paragraph_format
    fmt.alignment, fmt.line_spacing, fmt.space_before, fmt.space_after = JUSTIFY, 1.5, Pt(0), Pt(6)
    for name, size, bold, italic, caps, before, after in (
            ("Heading 1", 16, True, False, True, 0, 18), ("Heading 2", 14, True, True, False, 18, 12),
            ("Heading 3", 12, True, True, False, 12, 6)):
        style = doc.styles[name]
        style_font(style, "Times New Roman", size, bold, italic, caps)
        f = style.paragraph_format
        f.alignment, f.line_spacing, f.space_before, f.space_after = LEFT, 1.0, Pt(before), Pt(after)
        f.keep_with_next = True
    caption = paragraph_style(doc, "Caption")
    style_font(caption, "Times New Roman", 10)
    caption.paragraph_format.alignment, caption.paragraph_format.line_spacing = CENTER, 1.0
    caption.paragraph_format.space_after = Pt(6)
    front = paragraph_style(doc, "Naslov brez številke")
    style_font(front, "Times New Roman", 16, bold=True, caps=True)
    front.paragraph_format.alignment, front.paragraph_format.space_after = LEFT, Pt(18)
    small = paragraph_style(doc, "Naslov manjši")
    style_font(small, "Times New Roman", 12, bold=True, caps=True)
    small.paragraph_format.alignment, small.paragraph_format.space_after = LEFT, Pt(12)
    for level, (size, bold, italic, small_caps, indent) in {
            1: (12, True, False, False, 0.0), 2: (10, False, False, True, 0.6), 3: (10, False, True, False, 1.2)}.items():
        toc = paragraph_style(doc, f"toc {level}")
        style_font(toc, "Times New Roman", size, bold=bold, italic=italic, caps=level == 1, small_caps=small_caps)
        toc.paragraph_format.left_indent, toc.paragraph_format.line_spacing = Cm(indent), 1.15
        toc.paragraph_format.space_after = Pt(4 if level == 1 else 2)
    for name in ("List Bullet", "Footer"):
        style_font(doc.styles[name], "Times New Roman", 12)
    guide = paragraph_style(doc, "Navodilo")
    style_font(guide, "Times New Roman", 11, italic=True)
    guide.font.color.rgb = GUIDE
    guide.paragraph_format.alignment = LEFT
    reference = paragraph_style(doc, "Vir v seznamu")
    reference.paragraph_format.left_indent, reference.paragraph_format.first_line_indent = Cm(1.27), Cm(-1.27)
    reference.paragraph_format.alignment = LEFT


def guide(doc, text):
    return doc.add_paragraph(f"Navodilo: {text}", style="Navodilo")


def front_heading(doc, text, style="Naslov brez številke", new_page=True):
    paragraph = doc.add_paragraph(text, style=style)
    paragraph.paragraph_format.page_break_before = new_page
    return paragraph


def heading(doc, text, level, new_page=True):
    paragraph = doc.add_heading(text, level=level)
    if level == 1:
        paragraph.paragraph_format.page_break_before = new_page
    return paragraph


def caption(doc, kind, text):
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.add_run(f"{kind} ")
    add_field(paragraph, f"SEQ {kind} \\* ARABIC", "1")
    paragraph.add_run(f": {text}")
    return paragraph


def title_page(doc):
    def line(text, size=12, bold=False, before=0, after=0):
        paragraph = doc.add_paragraph()
        paragraph.alignment = CENTER
        paragraph.paragraph_format.space_before, paragraph.paragraph_format.space_after = Pt(before), Pt(after)
        paragraph.paragraph_format.line_spacing = 1.5
        run = paragraph.add_run(text)
        run.bold, run.font.size = bold, Pt(size)
        return paragraph

    line("VIŠJA STROKOVNA ŠOLA ACADEMIA", bold=True)
    line("MARIBOR", bold=True)
    line(TITLE.upper(), size=20, bold=True, before=150)
    line(f"Kandidat: {STUDENT}", before=130)
    for text in ("Vrsta študija: [študent rednega/izrednega študija]", "Študijski program: Računalništvo",
                 f"Mentor predavatelj: {MENTOR}", "Mentor v podjetju: [ime in priimek, izobrazba]",
                 "Lektor/-ica: [ime in priimek, izobrazba]"):
        line(text)
    line("Maribor, 2026", before=40)


def build_thesis(path: Path):
    doc = Document()
    thesis_styles(doc)
    page_setup(doc.sections[0], 3.0, 2.0, 2.5, 2.5)
    doc.core_properties.title, doc.core_properties.author = TITLE, STUDENT
    title_page(doc)

    front_heading(doc, "IZJAVA O AVTORSTVU DIPLOMSKEGA DELA")
    doc.add_paragraph(f"Podpisani {STUDENT} sem avtor diplomskega dela z naslovom {TITLE}, ki sem ga napisal pod "
                      "mentorstvom mag. Dušana Brgleza.")
    doc.add_paragraph("S svojim podpisom zagotavljam, da:")
    for text in (
            "je predloženo delo izključno rezultat mojega dela,",
            "sem poskrbel, da so dela in mnenja drugih avtorjev, ki jih uporabljam v predloženi nalogi, navedena oz. "
            "citirana skladno s pravili Višje strokovne šole Academia Maribor,",
            "se zavedam, da je plagiatorstvo – predstavljanje tujih del oz. misli kot mojih lastnih – kaznivo po Zakonu "
            "o avtorski in sorodnih pravicah (Uradni list RS, št. 16/07 – uradno prečiščeno besedilo, 68/08, 110/13, "
            "56/15 in 63/16 – ZKUASP); prekršek pa podleže tudi ukrepom Višje strokovne šole Academia Maribor skladno "
            "z njenimi pravili,",
            "skladno z 32.a členom ZASP dovoljujem Višji strokovni šoli Academia Maribor objavo diplomskega dela na "
            "spletnem portalu šole."):
        doc.add_paragraph(text, style="List Bullet")
    signature = doc.add_paragraph("Maribor, [mesec] 2026\t\t\t\t\tPodpis študenta:")
    signature.paragraph_format.space_before = Pt(48)
    guide(doc, "besedilo izjave je prevzeto iz diplomskih del Academia 2026; preverite ga z uradno predlogo v eAcademia.")

    front_heading(doc, "ZAHVALA", style="Naslov manjši")
    guide(doc, "neobvezno; zahvala mentorjema, podjetju in lektorici.")

    front_heading(doc, "POVZETEK", style="Naslov manjši")
    guide(doc, "približno ena stran (250–350 besed): problem, namen, metode, glavni rezultati po hipotezah in "
               "sklep. Povzetek napišite na koncu, ko so rezultati znani.")
    keywords = doc.add_paragraph()
    keywords.add_run("Ključne besede: ").bold = True
    keywords.add_run("optimizacija za iskalnike, Google Search Console, tehnični SEO, dostopnost, razlika razlik")

    front_heading(doc, "ABSTRACT", style="Naslov manjši")
    doc.add_paragraph().add_run(TITLE_EN).bold = True
    guide(doc, "angleški prevod povzetka.")
    keywords = doc.add_paragraph()
    keywords.add_run("Keywords: ").bold = True
    keywords.add_run("search engine optimization, Google Search Console, technical SEO, accessibility, "
                     "difference-in-differences")

    front_heading(doc, "KAZALO VSEBINE")
    add_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u',
              "Kazalo se posodobi ob odprtju dokumenta (desni klik → Posodobi polje).")
    for title, kind in (("KAZALO SLIK", "Slika"), ("KAZALO TABEL", "Tabela"), ("KAZALO GRAFIKONOV", "Grafikon")):
        front_heading(doc, title, new_page=title != "KAZALO GRAFIKONOV")
        add_field(doc.add_paragraph(), f'TOC \\h \\z \\c "{kind}"', "Posodobi polje.")

    front_heading(doc, "Uporabljeni angleški izrazi in kratice", style="Naslov manjši")
    for text in (
            "CLS: Cumulative Layout Shift (kumulativni premik postavitve)",
            "CTR: Click-Through Rate (razmerje med kliki in prikazi)",
            "CWV: Core Web Vitals (osrednji spletni vitalni kazalniki)",
            "DA: Domain Authority (metrika podjetja Moz)",
            "DiD: difference-in-differences (razlika razlik)",
            "E-E-A-T: Experience, Expertise, Authoritativeness, Trustworthiness",
            "EAA: European Accessibility Act (Direktiva (EU) 2019/882)",
            "GBP: Google Business Profile",
            "GSC: Google Search Console",
            "INP: Interaction to Next Paint",
            "LCP: Largest Contentful Paint",
            "PPC: Pay-Per-Click (plačano oglaševanje)",
            "SEO: Search Engine Optimization (optimizacija za iskalnike)",
            "SERP: Search Engine Results Page (stran z rezultati iskanja)",
            "WCAG: Web Content Accessibility Guidelines",
            "YMYL: Your Money or Your Life (vsebine, ki vplivajo na finance ali zdravje)"):
        doc.add_paragraph(text, style="List Bullet")

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.footer.is_linked_to_previous = False
    footer = body.footer.paragraphs[0]
    footer.alignment = CENTER
    add_field(footer, "PAGE", "1")

    heading(doc, "1 UVOD", 1, new_page=False)
    heading(doc, "1.1 Opis področja in opredelitev problema", 2)
    guide(doc, "osnutek iz popravljene dispozicije; razširite in dopolnite z viri.")
    for text in PROBLEM:
        doc.add_paragraph(text)
    heading(doc, "1.2 Namen, cilji in osnovne trditve", 2)
    doc.add_paragraph("Cilji diplomskega dela so:")
    for text in GOALS:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_paragraph("Raziskovalna vprašanja:")
    for text in RESEARCH_QUESTIONS:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_paragraph("Hipoteze:")
    for text in HYPOTHESES:
        doc.add_paragraph(text, style="List Bullet")
    heading(doc, "1.3 Predpostavke in omejitve", 2)
    guide(doc, "ena študija primera (ni posplošitve na vsa MSP); skupine strani niso naključne; posodobitve "
               "Googlovega algoritma in sezonskost; majhno število strani v skupini; Domain Authority ni metrika "
               "Googla; ZDPSI in izjema za mikropodjetja, ki opravljajo storitve; zaupnost podatkov podjetja.")
    heading(doc, "1.4 Uporabljene raziskovalne metode", 2)
    for text in METHODS:
        doc.add_paragraph(text)
    guide(doc, "za BTEC Unit 16 (M1, D1) utemeljite izbiro metod: pozitivistična filozofija, deduktivni pristop, "
               "kvantitativna metoda, kvazieksperiment, longitudinalni podatki (Saunders »Research Onion«), in "
               "kritično primerjajte z alternativami (anketa, intervju, A/B test).")

    chapters = {
        "2 TEORETIČNI DEL – SEO: OSNOVE IN RAZVOJ": {
            "2.1 Delovanje Googlovih algoritmov in dejavniki uvrstitve":
                "indeksiranje, PageRank (Brin in Page, 1998), sistemi za koristno vsebino, E-E-A-T in YMYL; empirične "
                "raziskave dejavnikov (Ziakis idr., 2019; Lewandowski idr., 2021).",
            "2.2 Tehnični SEO: hitrost, mobilnost in struktura":
                "Core Web Vitals (LCP, INP, CLS), mobile-first indeksiranje, kanonične povezave, statusne kode, "
                "najpogostejše napake po HTTP Archive (2024).",
            "2.3 Vsebinska optimizacija in ključne besede":
                "naslovi, meta opisi, H1, strukturirani podatki; odvisnost CTR od položaja (Craswell idr., 2008).",
            "2.4 Gradnja povratnih povezav (off-page SEO)":
                "vloga povezav, metrike Domain Authority in Domain Rating ter njihove omejitve (Moz, b. d.).",
            "2.5 Lokalni SEO in Google Business Profile": "dejavniki lokalnega uvrščanja (Whitespark, 2026).",
            "2.6 Dostopnost spletnih mest in SEO": "WCAG 2.1 in 2.2, EAA, ZDPSI; kje se dostopnost in SEO prekrivata.",
        },
        "3 PRAKTIČNI DEL – SEO-REVIZIJA IN OPTIMIZACIJA SPLETNEGA MESTA": {
            "3.1 Predstavitev izbranega spletnega mesta in začetno stanje":
                "spletno mesto, razdelitev strani v skupine, izhodiščne vrednosti (prikazi, kliki, položaj, Lighthouse).",
            "3.2 Izvedba SEO-revizije": "rezultati revizije po kategorijah napak; odgovor na RV4.",
            "3.3 Implementacija SEO-ukrepov in merjenje rezultatov":
                "dnevnik ukrepov z datumi, okna merjenja, zbiranje podatkov iz Google Search Console; RV2 in RV3.",
            "3.4 Dostopnost spletnega mesta (WCAG) in vpliv na SEO": "revizija z orodjema Lighthouse in WAVE; H4.",
            "3.5 Analiza rezultatov in robustnost":
                "ocene učinkov z intervali zaupanja, ablacija, placebo test, omejitve. Vsako številko prenesite iz "
                "report/claims.csv istega zagona.",
            "3.6 Odgovori na raziskovalna vprašanja in preverjanje hipotez":
                "za vsako RV in H en odstavek: ugotovitev, številka, odločitev (podprta / ni podprta / neodločeno).",
        },
    }
    for chapter, sections in chapters.items():
        heading(doc, chapter, 1)
        for number, (section, text) in enumerate(sections.items()):
            heading(doc, section, 2)
            guide(doc, text)
            if section.startswith("3.2"):
                figure = doc.add_paragraph("[slika]")
                figure.alignment = CENTER
                caption(doc, "Slika", "Število najdenih težav pred in po optimizaciji")
                source = doc.add_paragraph("(Vir: lasten)", style="Caption")
            if section.startswith("3.5"):
                caption(doc, "Tabela", "Ocene učinkov po hipotezah")
                table = doc.add_table(rows=2, cols=5)
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for cell, text_value in zip(table.rows[0].cells, ("Hipoteza", "Metrika", "Ocena", "Interval zaupanja",
                                                                  "Odločitev")):
                    cell.text = text_value
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            paragraph.alignment = LEFT
                            paragraph.paragraph_format.line_spacing = 1.0
                            for run in paragraph.runs:
                                run.font.size = Pt(10)
                doc.add_paragraph("Vir: lasten", style="Caption")

    heading(doc, "4 SKLEP", 1)
    guide(doc, "odgovori na namen in cilje, glavne ugotovitve, priporočila podjetju (BTEC D2), pomen za prakso, "
               "predlogi za nadaljnje raziskave. Brez novih rezultatov.")
    heading(doc, "4.1 Refleksija raziskovalnega procesa", 2)
    guide(doc, "BTEC Unit 16 (P6, P7, M4, D3): kako učinkovite so bile izbrane metode, katere alternativne metode bi "
               "bile mogoče, kaj bi naslednjič naredili drugače in zakaj.")

    heading(doc, "5 VIRI IN LITERATURA", 1)
    guide(doc, "APA 7 v slovenski obliki; abecedni vrstni red; [datum] zamenjajte z datumom dostopa. Seznam vsebuje "
               "preverjene vire; navedite le tiste, ki ste jih uporabili v besedilu, in dodajte slovenske vire.")
    for entry in REFERENCES:
        add_rich(doc.add_paragraph(style="Vir v seznamu"), entry)

    heading(doc, "6 PRILOGE", 1)
    for text in ("Priloga A: Protokol raziskave", "Priloga B: Dnevnik uvedenih SEO-ukrepov",
                 "Priloga C: Ponovljivost analize (koda, manifest zagona, kontrolne vsote)",
                 "Priloga D: Analiza deležnikov (BTEC Unit 16, LO3)"):
        doc.add_paragraph(text, style="List Bullet")

    update_fields_on_open(doc)
    doc.save(path)

# --------------------------------------------------------------------------- corrected dispozicija


def form_styles(doc):
    normal = doc.styles["Normal"]
    style_font(normal, "Arial", 10)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(3), 1.1
    style_font(doc.styles["List Bullet"], "Arial", 10)


def band(doc, text, fill="000000", color=RGBColor(255, 255, 255), size=11):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    shade(cell, fill)
    run = cell.paragraphs[0].add_run(text)
    run.bold, run.font.size, run.font.color.rgb = True, Pt(size), color
    return table


def block(doc, header, items):
    table = doc.add_table(rows=2, cols=1)
    table.style = "Table Grid"
    head, body = table.rows[0].cells[0], table.rows[1].cells[0]
    shade(head, "D9D9D9")
    head.paragraphs[0].add_run(header).bold = True
    first = True
    for kind, text in items:
        if kind == "li":
            paragraph = body.add_paragraph(style="List Bullet")
            add_rich(paragraph, text)
        else:
            paragraph = body.paragraphs[0] if first else body.add_paragraph()
            add_rich(paragraph, text, bold=kind == "h", italic=kind == "note")
            if kind == "ref":
                paragraph.paragraph_format.left_indent = Cm(0.8)
                paragraph.paragraph_format.first_line_indent = Cm(-0.8)
        first = False
    doc.add_paragraph()
    return table


def build_dispozicija(path: Path):
    doc = Document()
    form_styles(doc)
    page_setup(doc.sections[0], 2.2, 2.2, 2.0, 2.0)
    doc.core_properties.title, doc.core_properties.author = "Dispozicija – popravljena različica", STUDENT

    note = doc.add_paragraph()
    add_rich(note, "Popravljena različica dispozicije – predlog z dne 11. 9. 2026. Spremembe je treba uskladiti z "
                   "mentorjem; podpisana izvirna dispozicija z dne 29. 6. 2026 velja, dokler mentor sprememb ne potrdi. "
                   "Seznam popravkov je na koncu dokumenta.", italic=True, color=GUIDE)
    band(doc, "PREDLOGA ZA IZDELAVO DISPOZICIJE /\nPRIJAVA TEME RAZISKOVALNE NALOGE OZ. DIPLOMSKEGA DELA")
    doc.add_paragraph()
    info = doc.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    for label, value in (
            ("BTEC študijski program", "Pearson BTEC HND in Computing"),
            ("BTEC predmet", "Unit 16: Research Project (Set Theme 2026)"),
            ("SLO študijski program", "Računalništvo"), ("SLO predmet", "Diplomsko delo (DD)"),
            ("Mentor predavatelj (ime in priimek)", "mag. Dušan Brglez"),
            ("Mentor v podjetju (ime in priimek, izobrazba)", "[ime in priimek, izobrazba]"),
            ("Podjetje", "[naziv podjetja – v javni različici le ob soglasju za objavo]"),
            ("Datum objave", "22. 5. 2026"), ("Rok za oddajo dispozicije DD", "30. 6. 2026"),
            ("Rok za oddajo DD", "30. 12. 2026")):
        cells = info.add_row().cells
        shade(cells[0], "D9D9D9")
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].text = value
    doc.add_paragraph()
    for label, value in (("Ime in priimek študenta:", STUDENT), ("Vpisna številka:", "[vpisna številka]"),
                         ("Predlagani naslov DD:", TITLE), ("Datum:", "[datum]")):
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{label} ").bold = True
        paragraph.add_run(value)
    doc.add_paragraph()

    block(doc, "1. POGLAVJE: STROKOVNO PODROČJE, NAMEN, OPREDELITEV RAZISKOVALNEGA PROBLEMA, RAZISKOVALNA "
               "VPRAŠANJA, HIPOTEZE, CILJI, KROVNA TEMA",
          [("h", "Strokovno področje, namen in opredelitev raziskovalnega problema:")]
          + [("p", text) for text in PROBLEM]
          + [("h", "Cilji:")] + [("li", text) for text in GOALS]
          + [("h", "Raziskovalna vprašanja raziskovalne naloge oz. diplomskega dela:")]
          + [("li", text) for text in RESEARCH_QUESTIONS]
          + [("h", "Hipoteze / trditve diplomskega dela:")] + [("li", text) for text in HYPOTHESES]
          + [("h", "Raziskovalni proces in metode:")] + [("p", text) for text in METHODS]
          + [("h", "Izdelek:"), ("p", PRODUCT)])

    outline = []
    for chapter, sections in OUTLINE:
        outline.append(("h", chapter))
        outline += [("ref", section) for section in sections]
    block(doc, "2. POGLAVJE: PREDVIDENO KAZALO", outline)

    block(doc, "3. POGLAVJE: LITERATURA IN VIRI",
          [("note", "Predvidena literatura in viri (vsaj 8 virov, od tega vsaj 1 tuji vir); slovenske strokovne vire "
                    "iz COBISS je treba še dodati.")]
          + [("ref", f"{number}. {entry}") for number, entry in enumerate(REFERENCES, start=1)])

    block(doc, "4. POGLAVJE: AKTIVNOSTI IN ČASOVNI NAČRT", [
        ("li", "Pregled literature (september–oktober 2026)."),
        ("li", "Dostopi do Google Search Console, Google Analytics in Google Business Profile; razdelitev strani v "
               "skupine in zapis protokola raziskave (do 20. 9. 2026)."),
        ("li", "SEO-revizija (Screaming Frog, Google Search Console, Lighthouse, WAVE) in izvoz podatkov predobdobja; "
               "Google Search Console hrani podatke za 16 mesecev (do 25. 9. 2026)."),
        ("li", "Uvedba ukrepov po skupinah strani z zamaknjenimi datumi (21. 9.–5. 10. 2026)."),
        ("li", "Optimizacija Google Business Profile in gradnja povratnih povezav (od 1. 10. 2026)."),
        ("li", "Merjenje: 2 tedna uvajanja in 6 tednov poobdobja po zadnji skupini (do 30. 11. 2026)."),
        ("li", "Analiza podatkov in preverjanje hipotez (3.–10. 12. 2026)."),
        ("li", "Pisanje teoretičnega dela (september–oktober), praktičnega dela in sklepa (november–december 2026)."),
        ("li", "Lektoriranje, pregled z mentorjem in oddaja (do 30. 12. 2026)."),
        ("h", "Kritične točke in mejniki:"),
        ("p", "Kritična točka 1: zaključek teoretičnega dela in SEO-revizije – oddaja 1. in 2. poglavja mentorju. "
              "Predvideni datum (dogovorjen z mentorjem): [predlog 20. 10. 2026]"),
        ("p", "Kritična točka 2: zaključek merjenja in analize – oddaja 3. poglavja mentorju. Predvideni datum "
              "(dogovorjen z mentorjem): [predlog 15. 12. 2026]"),
        ("p", "Terminski načrt: september 2026 – literatura, dostopi in revizija; konec septembra in začetek oktobra – "
              "uvedba ukrepov; oktober–november – merjenje; december – analiza, pisanje, lektoriranje in oddaja."),
    ])

    for header, lines in (
            ("ŠTUDENT", ["Potrjujem, da je ta rezultat mojega lastnega dela.",
                         "Raziskovalna naloga oz. diplomsko delo ni bilo ali ne bo uporabljeno za druge namene.",
                         "Datum: ____________    Ime in priimek ter podpis: ______________________"]),
            ("MENTOR NA ŠOLI", ["Potrjujem dispozicijo raziskovalne naloge / diplomskega dela in mentorstvo.",
                                "Datum: ____________    Ime in priimek ter podpis: ______________________"]),
            ("MENTOR V PODJETJU", ["Potrjujem dispozicijo raziskovalne naloge / diplomskega dela in mentorstvo v podjetju.",
                                   "Datum: ____________    Ime in priimek ter podpis: ______________________"]),
            ("IZJAVA PODJETJA", ["Soglašamo, da lahko študent/-ka na primeru našega podjetja izdela raziskovalno "
                                 "nalogo / diplomsko delo.",
                                 "Dovoljujemo naziv podjetja v raziskovalni nalogi / diplomskem delu (ustrezno obkrožite): "
                                 "objavo dovolimo / objave ne dovolimo",
                                 "Datum: ____________    Ime in priimek ter podpis odgovorne osebe: ______________  žig:"])):
        block(doc, header, [("p", line) for line in lines])

    doc.add_page_break()
    band(doc, "SEZNAM POPRAVKOV GLEDE NA IZVIRNO DISPOZICIJO (29. 6. 2026)", fill="D9D9D9", color=BLACK, size=10)
    corrections = [
        ("Opredelitev problema", "Google ima v Sloveniji »več kot 95 %« trga",
         "približno 93 % (StatCounter, avgust 2026: 93,2 %)", "podatek brez vira in ni točen"),
        ("Opredelitev problema", "»primanjkljaj empiričnih analiz« brez vira",
         "previdnejša trditev s sklicem na literaturo", "trditve o vrzeli ni bilo mogoče podpreti"),
        ("RV1", "le tehnični SEO, vsebina, povezave", "vse štiri skupine ukrepov iz hipotez",
         "usklajenost vprašanj in hipotez"),
        ("H1", "»izboljšanje uvrstitve za vsaj 20 % v primerjavi z začetnim stanjem«",
         "povprečni položaj, utežen s prikazi, glede na kontrolne strani; 6 tednov",
         "metrika ni bila določena; primerjava z začetnim stanjem ne loči učinka od sezonskosti in posodobitev algoritma"),
        ("H2", "CTR brez upoštevanja položaja", "CTR, prilagojen položaju, glede na kontrolne strani",
         "CTR je močno odvisen od položaja"),
        ("H3", "»za slovenska mala in srednja podjetja«", "»na izbranem spletnem mestu«; primerjava z vsebinsko optimizacijo",
         "ena študija primera ne omogoča posplošitve"),
        ("H4", "»uporabnišku«; »iskalniki nagrajujejo dostopna spletna mesta« kot dejstvo",
         "»uporabniško«; ničelni izid je dopusten", "tipkarska napaka; Google dostopnosti ne navaja kot dejavnika uvrščanja"),
        ("H5", "+10 točk Domain Authority v 6 mesecih kot vzročna trditev", "opisna trditev v obdobju raziskave",
         "6 mesecev presega rok oddaje; DA je metrika Moz; ni kontrolne skupine"),
        ("Metode", "primerjava pred/po", "kvazieksperiment s kontrolno skupino, razlika razlik, bootstrap, Bonferroni, placebo",
         "zanesljivejša ocena učinka"),
        ("Izdelek", "revizija in akcijski načrt", "dodano ponovljivo orodje v Pythonu, predstavitvena stran in poročilo",
         "izdelek za podjetje in BTEC Unit 16 (LO3)"),
        ("Kazalo", "4 Sklep, 5 Literatura in viri", "dodani 2.6, 3.5, 3.6, 4.1 Refleksija in 6 Priloge",
         "struktura diplomskih del Academia 2026 in merila BTEC Unit 16 (P6, P7, M4, D3)"),
        ("Literatura", "Fishkin, R. (2012). The Art of SEO (3rd ed.)", "Enge, Spencer in Stricchiola (2023), 4. izdaja",
         "izdaja iz leta 2012 je 2. izdaja s prvim avtorjem Enge; 3. izdaja je izšla 2015"),
        ("Literatura", "Searchmetrics (2024). Ranking Factors Study", "znanstveni članki (Ziakis idr., 2019; Lewandowski idr., 2021)",
         "študije ni bilo mogoče najti; Searchmetrics je od 2023 del podjetja Conductor"),
        ("Literatura", "pretežno blogi podjetij", "znanstveni članki, standardi, zakonodaja RS, APA 7",
         "zahteve za diplomsko delo"),
        ("Oštevilčenje", "poglavja 1, 2, 3, 5", "poglavja 1, 2, 3, 4", "manjkajoče 4. poglavje"),
        ("Časovni načrt", "junij–september 2026", "september–december 2026", "prvotni roki so pretekli; rok oddaje 30. 12. 2026"),
    ]
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Št.", "Mesto", "Prvotno", "Popravek", "Razlog")):
        shade(cell, "D9D9D9")
        cell.paragraphs[0].add_run(text).bold = True
    for number, row in enumerate(corrections, start=1):
        cells = table.add_row().cells
        for cell, text in zip(cells, (str(number), *row)):
            cell.text = text
    table.autofit = False
    widths = (Cm(0.8), Cm(2.6), Cm(4.5), Cm(4.5), Cm(4.2))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
            for paragraph in cell.paragraphs:
                paragraph.alignment = LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    # A document cannot end with a table; a full-size trailing paragraph would spill onto an empty page.
    tail_style = paragraph_style(doc, "Konec dokumenta")
    style_font(tail_style, "Arial", 1)
    tail = doc.add_paragraph(style=tail_style)
    tail.paragraph_format.space_before = tail.paragraph_format.space_after = Pt(0)
    tail.paragraph_format.line_spacing = Pt(1)
    doc.save(path)


def main():
    thesis = THESIS_DIR / "Diplomsko_delo_predloga_Rakhmanov.docx"
    dispozicija = THESIS_DIR / "Dispozicija_popravljena_Rakhmanov.docx"
    build_thesis(thesis)
    build_dispozicija(dispozicija)
    print(f"wrote {thesis}\nwrote {dispozicija}")


if __name__ == "__main__":
    main()
